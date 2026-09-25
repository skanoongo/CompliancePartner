#!/usr/bin/env python3
"""
NetSuite Change Management (SOX CM-02) - customization change evidence.

READ-ONLY. Opens NetSuite's customization lists, screenshots them, and reads the
rows. It never creates, edits, deploys or deletes a script, workflow or record type.

WHAT THIS CAPTURES, AND WHAT IT DOES NOT
----------------------------------------
Change management for NetSuite means: which customizations exist, and when each
one last changed. This walks the customization lists - scripts, deployments,
workflows, custom record types, custom fields, forms - and records the columns
NetSuite shows, which include the owner and, on most lists, a date.

It does NOT reconstruct an approval trail. NetSuite's per-record System Notes hold
who changed what, but only one record at a time and only for a bounded window, and
approvals live in whatever ticketing system the change was raised in. So this
produces the POPULATION of customizations and their last-changed dates, which is
the part a script can evidence honestly. Matching that population to approved
tickets remains a human step, and the workbook says so rather than implying the
control is satisfied by the listing alone.

The period is a label. These lists show current state, not a historical view, so a
run "for Q3 FY26" evidences the customizations as they stood when it ran.

Usage
-----
  python3 -m netsuite.sox_capture --period "Q3 FY26"
  python3 -m netsuite.sox_capture --area scripts,workflows
"""

import argparse
import csv
import json
import math
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from playwright.sync_api import sync_playwright

from core.platform import (
    PX_PER_ROW,
    launch_browser,
    log,
    now_stamp,
    phase,
    slug,
    workbook_copy,
)
from core import environments
from netsuite import common as ns

# The customization areas a change review covers. Each becomes a workbook tab and
# is selectable with --area.
AREAS = [
    {"key": "scripts", "label": "Scripts",
     "path": "/app/common/scripting/scripts.nl?whence="},
    {"key": "deployments", "label": "Script deployments",
     "path": "/app/common/scripting/scriptdeployments.nl?whence="},
    {"key": "workflows", "label": "Workflows",
     "path": "/app/common/workflow/setup/workflowlist.nl?whence="},
    {"key": "recordtypes", "label": "Custom record types",
     "path": "/app/common/custom/custrecords.nl?whence="},
    {"key": "fields", "label": "Custom fields (entity)",
     "path": "/app/common/custom/entitycustfields.nl?whence="},
    {"key": "bodyfields", "label": "Custom fields (transaction body)",
     "path": "/app/common/custom/bodycustfields.nl?whence="},
    {"key": "forms", "label": "Custom forms (transaction)",
     "path": "/app/common/custom/custforms.nl?whence="},
]


def area_keys():
    return [a["key"] for a in AREAS]


def select_areas(selection):
    """Narrow AREAS to the chosen ones. An unknown name is an error, never a
    silent fall back to everything - a review scoped to scripts that quietly
    captured all seven would be wrong in a way nobody would notice."""
    if not selection or selection.strip().lower() in ("", "all", "*"):
        return AREAS
    wanted = [w.strip().lower() for w in selection.split(",") if w.strip()]
    chosen, missing = [], []
    for w in wanted:
        hit = [a for a in AREAS if a["key"] == w or w in a["label"].lower()]
        if not hit:
            missing.append(w)
        for a in hit:
            if a not in chosen:
                chosen.append(a)
    if missing:
        raise SystemExit("No such area: " + ", ".join(missing)
                         + "\nKnown areas: " + " | ".join(area_keys()))
    return chosen


def list_verdict(data):
    if not data or not data.get("rows"):
        return False, "no table rows"
    headers = [h.lower() for h in data.get("headers", [])]
    wanted = ("name", "id", "script", "type", "owner", "date", "status", "record",
              "label", "form", "workflow")
    hits = [h for h in headers if any(w in h for w in wanted)]
    if hits:
        return True, f"columns {hits[:4]}"
    if len(data["rows"]) >= 3:
        return True, f"{len(data['rows'])} rows, unnamed columns"
    return False, f"{len(data['rows'])} rows and no recognisable columns"


def collect(cap, areas):
    """Open each area, screenshot it, read the rows."""
    out = []
    base = ns.account_base(cap.account_id)

    for area in areas:
        url = base + area["path"]
        log(f"  {area['label']}: {url}")
        if not cap.goto(url):
            cap.warn(f"could not open {area['label']} at {url} - the signed-in role may "
                     f"not have permission for it")
            out.append({**area, "url": url, "headers": [], "rows": [], "read": False})
            continue

        data = ns.read_list(cap.page)
        ok, why = list_verdict(data)
        cap.shoot_scrolled("changes", f"{area['label']} - {cap.account_id}", url)
        if not ok:
            cap.warn(f"{area['label']}: {why} - screenshot kept, but no rows were read")
            out.append({**area, "url": url, "headers": [], "rows": [], "read": False})
            continue

        log(f"    {len(data['rows'])} row(s) - {why}")
        out.append({**area, "url": url, "headers": data["headers"],
                    "rows": data["rows"], "read": True})
    return out


# --------------------------------------------------------------------------- #
HEAD_FILL = PatternFill("solid", fgColor="1D2433")


def write_outputs(out_dir, areas, meta):
    (out_dir / "changes.json").write_text(json.dumps({**meta, "areas": areas}, indent=2))
    with (out_dir / "changes.csv").open("w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["area", "column_1", "column_2", "column_3", "column_4",
                    "column_5", "column_6", "full_row"])
        for a in areas:
            for row in a["rows"]:
                w.writerow([a["label"]] + (row + [""] * 6)[:6] + [" | ".join(row)[:500]])
    total = sum(len(a["rows"]) for a in areas)
    log(f"  wrote changes.csv and changes.json ({total} rows across "
        f"{sum(1 for a in areas if a['read'])} list(s))")


def build_workbook(manifest_path, out_path, max_width=1400):
    data = json.loads(Path(manifest_path).read_text())
    areas = data["areas"]
    meta = {k: data.get(k, "") for k in ("period", "account", "captured")}

    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet(title="Summary")
    ws["A1"] = f"NetSuite Change Management - {meta['period']}"
    ws["A1"].font = Font(bold=True, size=16)
    rows = [("Account", meta["account"]), ("Listed", meta["captured"]), ("", "")]
    for a in areas:
        rows.append((a["label"], len(a["rows"]) if a["read"] else "not readable"))
    for i, (k, v) in enumerate(rows, start=3):
        ws.cell(row=i, column=1, value=k).font = Font(bold=bool(k))
        ws.cell(row=i, column=2, value=v)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 50

    n = len(rows) + 4
    ws.cell(row=n, column=1, value=(
        "This is the POPULATION of customizations and the columns NetSuite shows for "
        "them, captured as it stood when the run executed and labelled "
        f"{meta['period']}. It is not an approval trail: NetSuite keeps per-record "
        "System Notes one record at a time, and approvals live in the ticketing system "
        "the change was raised in. Matching this population to approved changes is a "
        "human step and is NOT evidenced by this workbook."))
    ws.cell(row=n, column=1).alignment = Alignment(wrap_text=True)
    ws.row_dimensions[n].height = 72

    for a in areas:
        sheet = wb.create_sheet(title=a["label"][:31])
        sheet["A1"] = f"{a['label']} - {meta['period']}"
        sheet["A1"].font = Font(bold=True, size=14)
        sheet["A2"] = a["url"]
        if not a["read"]:
            sheet["A4"] = ("This list could not be read. See the screenshots and the "
                           "warnings - this is not evidence that it is empty.")
            sheet["A4"].font = Font(bold=True, color="98600C")
            continue
        headers = a["headers"] or [f"Column {i+1}" for i in range(
            max((len(r) for r in a["rows"]), default=1))]
        for i, h in enumerate(headers, start=1):
            c = sheet.cell(row=4, column=i, value=h)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = HEAD_FILL
            sheet.column_dimensions[c.column_letter].width = 28
        for r, row in enumerate(a["rows"], start=5):
            for i, v in enumerate(row, start=1):
                sheet.cell(row=r, column=i, value=v)

    ev = wb.create_sheet(title="Evidence")
    ev.column_dimensions["A"].width = 200
    ev["A1"] = f"Screenshots - {meta['period']}"
    ev["A1"].font = Font(bold=True, size=14)
    img_dir = Path(manifest_path).parent / "_workbook_images"
    r = 3
    for c in data.get("captures", []):
        ev.cell(row=r, column=1, value=c["description"]).font = Font(bold=True)
        ev.cell(row=r + 1, column=1, value=c["url"])
        ts = datetime.fromisoformat(c["captured"]).strftime("%Y-%m-%d %H:%M:%S %Z")
        ev.cell(row=r + 2, column=1, value=f"Captured: {ts}")
        img = XLImage(workbook_copy(c["file"], img_dir))
        if img.width > max_width:
            scale = max_width / img.width
            img.width = int(img.width * scale)
            img.height = int(img.height * scale)
        ev.add_image(img, f"A{r + 3}")
        r = r + 3 + math.ceil(img.height / PX_PER_ROW) + 3
    if r == 3:
        ev["A3"] = "No screenshots were captured."

    wb.save(out_path)
    log(f"Workbook written: {out_path}")


def run(args, out_dir, areas):
    with sync_playwright() as p:
        ctx = launch_browser(p, args.profile, args.browser)
        page = ctx.pages[0] if ctx.pages else ctx.new_page()
        page.set_default_timeout(60000)

        page = ns.sign_in(ctx, page, args.account, args.username, args.password,
                          args.login, args.session_url, role=args.role)
        if page is None:
            log("Sign-in did not complete. Exiting.")
            try:
                ctx.close()
            except Exception:
                pass
            sys.exit(1)
        page.set_default_timeout(60000)
        page.bring_to_front()
        log(f"Signed in. Current URL: {page.url}")

        if len(areas) != len(AREAS):
            log(f"Scoped to {len(areas)} of {len(AREAS)} areas: "
                + ", ".join(a["label"] for a in areas))
        phase("capturing", str(len(areas)))

        cap = ns.NsCapture(page, out_dir, args.settle, args.account, args.display,
                           args.max_parts)
        collected = collect(cap, areas)

        meta = {"period": args.period, "account": args.account,
                "captured": now_stamp().isoformat()}
        write_outputs(out_dir, collected, meta)

        manifest = out_dir / "manifest.json"
        manifest.write_text(json.dumps({
            **meta, "areas": collected, "captures": cap.manifest,
            "warnings": cap.warnings}, indent=2))
        log(f"Manifest written: {manifest} ({len(cap.manifest)} screenshots)")

        if not any(a["read"] for a in collected):
            ctx.close()
            if cap.auth_lost:
                raise SystemExit(
                    "The NetSuite session was not valid: every list redirected back to a "
                    "sign-in page.\nThis is NOT a permissions problem - the sign-in never "
                    "completed, usually because a two-factor code is still outstanding.\n"
                    "Open the browser session, finish signing in, and re-run.")
            raise SystemExit(
                "None of the customization lists could be read, but the session was "
                "valid - so the signed-in role most likely lacks permission for "
                "Customization > Scripting. Switch to an administrator role and re-run.")

        if cap.warnings:
            print("\n==== WARNINGS - resolve these before signing off ====")
            for w in cap.warnings:
                print("  " + w)
            print()
        ctx.close()
        return manifest


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--period", default="")
    ap.add_argument("--account", default="")
    ap.add_argument("--out", default=None)
    ap.add_argument("--profile", default="~/.netsuite_sox_browser_profile")
    ap.add_argument("--browser", choices=["chromium", "chrome"], default="chromium")
    ap.add_argument("--settle", type=float, default=4.0)
    ap.add_argument("--display", type=int, default=1)
    ap.add_argument("--max-parts", type=int, default=12)
    ap.add_argument("--area", default="all",
                    help="areas to capture, comma-separated, or all (" +
                         " | ".join(area_keys()) + ")")
    ap.add_argument("--list-areas", action="store_true")
    ap.add_argument("--capture-only", action="store_true")
    ap.add_argument("--build-only", action="store_true")
    ap.add_argument("--no-pause", action="store_true")
    ap.add_argument("--env", default="netsuite")
    ap.add_argument("--username", default="")
    ap.add_argument("--role", default="",
                    help="role to select on the NetSuite role picker, e.g. Administrator. "
                         "Only chosen when it matches BOTH the role name and the account.")
    args = ap.parse_args()

    if args.list_areas:
        for a in AREAS:
            print(f"{a['key']:14} {a['label']}")
        return

    areas = select_areas(args.area)

    if not args.period:
        args.period = f"Q{(datetime.now().month - 1) // 3 + 1} {datetime.now().year}"

    args.password = ""
    args.login = "interactive"
    args.role = args.role or ""
    args.session_url = __import__("os").environ.get("CAPTURE_SESSION_URL", "")
    try:
        cfg = environments.find(args.env)
        if cfg and cfg.get("enabled"):
            args.username = args.username or cfg.get("username", "")
            args.password = cfg.get("password", "")
            args.login = cfg.get("login", "interactive")
            args.account = args.account or cfg.get("account_id", "")
            args.role = args.role or cfg.get("role", "Administrator")
            log(f"Environment '{cfg['name']}': login={args.login}"
                f"{', user=' + args.username if args.username else ''}"
                f"{', password configured' if args.password else ', no password configured'}")
    except Exception as exc:
        log(f"!! environments.yaml could not be read ({exc}); signing in by hand.")

    if not args.account:
        sys.exit("No NetSuite account id. Set account_id in config/environments.yaml "
                 "or pass --account 7258820_SB1.")

    out_dir = Path(args.out) if args.out else Path(
        f"ns_cm_{slug(args.period)}_{datetime.now().strftime('%Y%m%d-%H%M%S')}")
    out_dir.mkdir(parents=True, exist_ok=True)
    xlsx = out_dir / f"NetSuite Change Management - {args.period}.xlsx"

    if args.build_only:
        m = out_dir / "manifest.json"
        if not m.exists():
            sys.exit(f"No manifest at {m}")
        build_workbook(m, xlsx)
        return

    phase("starting", args.period)
    manifest = run(args, out_dir, areas)
    if not args.capture_only:
        phase("building_workbook", str(xlsx))
        build_workbook(manifest, xlsx)
    phase("done", str(xlsx))
    log("Done.")


if __name__ == "__main__":
    main()
