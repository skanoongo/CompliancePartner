#!/usr/bin/env python3
"""
Workato User Access Review (SOX UA-04) - collaborator listing + screenshot evidence.

READ-ONLY. It opens the workspace's collaborator pages, reads who is listed, and
screenshots what it read. It never invites, removes, suspends or edits anyone.

WHAT IT PRODUCES
----------------
  users.csv / users.json   every collaborator: name, email, role, status, whether
                           this review counts them as active, and - when a Workday
                           terminations list is supplied - whether they left
  screenshots + manifest   whole-screen captures of the pages the list came from,
                           so the listing can be tied back to what was on screen
  .xlsx workbook           the UAR template's five tabs: Review Checklist |
                           Parameter Screenshot - Before | User List | Workday
                           termination List | Parameter Screenshot - After

WHAT IT READS
-------------
  - the collaborator roster, at the first candidate URL that really is one, every
    page of it: scrolled for lists that render on demand, and "next" followed for
    lists that paginate
  - each collaborator's own page, for the role they hold in THIS workspace's
    environment - the list shows one role, the detail page shows it per
    environment, and the review is of one environment (--no-detail skips this)
  - any further access pages configured as `uar_pages` (Workato Agentic users
    and roles, say), each listed and evidenced the same way, tagged by source

TERMINATIONS
------------
The Workday "SOX Audit - Terminations" export (.csv or .xlsx), from --terminations
or the Workato environment's `uar_terminations` setting (Doppler:
ENVIRONMENTS_WORKATO_UAR_TERMINATIONS), is matched by email. A terminated person
still on the roster is flagged in the workbook, in users.json and as a run
warning. Without the file the column says "not checked" - never "Active User",
which would be a claim nobody verified.

ON "ACTIVE" AND "INACTIVE"
--------------------------
Workato does not label collaborators with a single active/inactive flag. It shows a
state per row - active, pending/invited, suspended, deactivated - and the wording
varies by plan and by page. This script records the RAW status exactly as the page
showed it, and separately derives an `active` verdict from it. Both go in the
output. When a status is one it does not recognise, the row is marked `unknown` and
raised as a warning rather than quietly counted as active: a user access review that
under-counts access is the failure that matters.

POINT IN TIME
-------------
The listing is as-of the moment it runs. Workato does not expose a historical
collaborator list, so a review "for Q3" is a current snapshot labelled with that
period, evidenced by the timestamp in each screenshot. That is how the period
argument is used, and the workbook says so on its face.

Usage
-----
  python3 -m workato.uar_capture --period "Q3 FY26"
  python3 -m workato.uar_capture --period "Q3 FY26" --out ./uar_q3
  python3 -m workato.uar_capture --capture-only      # no workbook
  python3 -m workato.uar_capture --terminations /data/terminations.xlsx
  python3 -m workato.uar_capture --no-detail         # roster only, no per-user pages
"""

import argparse
import csv
import json
import math
import re
import sys
import time
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from playwright.sync_api import sync_playwright

from core.platform import (
    APP_NAME,
    PX_PER_ROW,
    SCROLL_BY_JS,
    SCROLL_JS,
    activate_app,
    grab_screen,
    interactive,
    launch_browser,
    log,
    now_stamp,
    phase,
    slug,
    workbook_copy,
)
from workato.sox_capture import (
    BASE,
    ensure_workspace,
    page_text,
)
from core import environments

# Where collaborators live. Workato has moved this more than once and it differs by
# plan, so every candidate is tried and the first one that actually renders a list
# of people wins. The one that worked is recorded in the manifest.
MEMBER_URLS = (
    # Current: Workspace admin > Access control > Collaborators. With environments
    # enabled it exists only in the Development environment, and only for an
    # account allowed to manage collaborators - /members/collaborators redirects
    # here, and anyone else gets Workato's "That page doesn't exist".
    f"{BASE}/members/access_control/collaborators",
    f"{BASE}/members/access_control",
    f"{BASE}/members/collaborators",
    f"{BASE}/collaborators",
    f"{BASE}/members/users",
    f"{BASE}/members/team",
    f"{BASE}/settings/collaborators",
    f"{BASE}/workspace/collaborators",
    f"{BASE}/settings/members",
    f"{BASE}/admin/members",
    f"{BASE}/teams",
    # Last: bare /members redirects to whichever admin sub-page Workato defaults to
    # (error alerts, in at least one workspace), which is not a roster at all.
    f"{BASE}/members",
)

# A page has to look like a roster, not merely contain an email address. The first
# version of this checked only "is there an email on the page", and happily accepted
# the Workspace admin error-alerts screen, reading a notification recipient chip as
# the entire collaborator list. Text on a real roster page.
ROSTER_WORDS = ("collaborator", "member", "role", "invite", "permission",
                "last active", "team", "access")
# Column headers a roster has and a settings form does not.
ROSTER_HEADERS = ("email", "role", "status", "name", "collaborator", "member",
                  "user", "last active", "permission")

EMAIL_RE = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")

# Raw status -> does this review count the person as holding access?
# Anything not listed here is "unknown" and gets a warning; it is never assumed
# inactive, because assuming inactive is what hides live access from a reviewer.
STATUS_ACTIVE = ("active", "enabled", "accepted", "member", "admin", "analyst", "operator")
STATUS_INACTIVE = ("deactivated", "disabled", "suspended", "removed", "revoked", "inactive", "expired")
STATUS_PENDING = ("pending", "invited", "invitation sent", "awaiting")

# Pull every row that mentions an email address, from a table or from the card and
# list layouts Workato uses on narrower pages. Scrolling re-runs this, and rows are
# merged by email, so a virtualised list that only renders what is on screen still
# ends up complete.
EXTRACT_JS = r"""
() => {
  const EMAIL = /[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}/;
  const out = [];
  const seen = new Set();

  // An email inside a form control is a setting, not a person on a roster - the
  // notification-recipient chip on Workato's admin screen is exactly this, and
  // reading it as the collaborator list is how this went wrong before.
  const inFormControl = (el) => {
    for (let n = el; n; n = n.parentElement) {
      const tag = n.tagName;
      if (tag === 'FORM' || tag === 'INPUT' || tag === 'TEXTAREA' || tag === 'SELECT') return true;
      if (n.getAttribute && n.getAttribute('contenteditable') === 'true') return true;
      const role = n.getAttribute && n.getAttribute('role');
      if (role === 'combobox' || role === 'textbox' || role === 'listbox') return true;
    }
    return false;
  };

  const push = (cells, el, source) => {
    if (el && inFormControl(el)) return;
    const text = cells.filter(Boolean).join(' | ');
    const m = text.match(EMAIL);
    if (!m) return;
    const key = m[0].toLowerCase();
    if (seen.has(key)) return;
    seen.add(key);
    // The collaborator's own page, when the row links to it - that is where the
    // per-environment role is.
    const a = el && el.querySelector && el.querySelector('a[href*="/members/"]');
    out.push({email: m[0], cells: cells, text: text, source: source,
              href: a ? a.getAttribute('href') : ''});
  };

  let fromTable = 0;
  document.querySelectorAll('tr').forEach(tr => {
    const cells = Array.from(tr.querySelectorAll('td,th')).map(c => (c.innerText||'').trim());
    if (!cells.length) return;
    const before = out.length;
    push(cells, tr, 'table');
    if (out.length > before) fromTable++;
  });

  if (out.length === 0) {
    // No table: the smallest elements holding exactly one email, so a card or list
    // row is captured without dragging the whole page in with it.
    const candidates = Array.from(document.querySelectorAll('li,div,article,section'))
      .filter(el => {
        if (inFormControl(el)) return false;
        const t = (el.innerText || '').trim();
        if (!t || t.length > 400) return false;
        return (t.match(new RegExp(EMAIL.source, 'g')) || []).length === 1;
      });
    candidates.sort((a, b) => (a.innerText||'').length - (b.innerText||'').length);
    const claimed = [];
    candidates.forEach(el => {
      if (claimed.some(c => c.contains(el) || el.contains(c))) return;
      claimed.push(el);
      const lines = (el.innerText || '').split('\n').map(s => s.trim()).filter(Boolean);
      push(lines, el, 'card');
    });
  }

  const headers = Array.from(document.querySelectorAll('thead th, tr th'))
    .map(th => (th.innerText || '').trim()).filter(Boolean);

  const body = (document.body && document.body.innerText || '').toLowerCase();
  // Every email on the page, including ones in form controls, purely so the caller
  // can tell "a page with people on it" from "a settings page with one address".
  const allEmails = (document.body && document.body.innerText || '')
    .match(new RegExp(EMAIL.source, 'g')) || [];

  return {
    rows: out, headers: headers, url: location.href, title: document.title,
    fromTable: fromTable,
    emailCount: new Set(allEmails.map(e => e.toLowerCase())).size,
    bodyText: body.slice(0, 4000),
  };
}
"""


def classify(raw):
    """Map a raw status string onto the verdict this review uses.

    Returns (active, bucket). `active` is None when the status is not recognised -
    deliberately not False, so an unknown state can never be silently written off
    as "no access".
    """
    t = (raw or "").strip().lower()
    if not t:
        return None, "unknown"
    for s in STATUS_INACTIVE:
        if s in t:
            return False, "inactive"
    for s in STATUS_PENDING:
        if s in t:
            # Pending access is not yet access, but it is not nothing either - it is
            # reported in its own bucket so a reviewer decides rather than the tool.
            return False, "pending"
    for s in STATUS_ACTIVE:
        if s in t:
            return True, "active"
    return None, "unknown"


def status_from_row(cells, headers):
    """Best guess at which cell holds the status, and what it says."""
    lower_headers = [h.lower() for h in headers]
    for i, h in enumerate(lower_headers):
        if "status" in h or "state" in h:
            if i < len(cells) and cells[i].strip():
                return cells[i].strip()
    # No status column: take the first cell whose text is a status word.
    for c in cells:
        t = c.strip().lower()
        if any(s in t for s in STATUS_INACTIVE + STATUS_PENDING + ("active",)):
            return c.strip()
    return ""


def field_from_row(cells, headers, names):
    for i, h in enumerate([h.lower() for h in headers]):
        if any(n in h for n in names):
            if i < len(cells) and cells[i].strip():
                return cells[i].strip()
    return ""


def name_from_row(cells, headers, email):
    got = field_from_row(cells, headers, ("name", "user", "collaborator", "member"))
    if got and "@" not in got:
        return got
    for c in cells:
        t = c.strip()
        if t and "@" not in t and not EMAIL_RE.search(t) and len(t) < 80:
            if not any(s in t.lower() for s in STATUS_INACTIVE + STATUS_PENDING + ("active",)):
                return t
    return email.split("@")[0]


MEMBER_LINK_RE = re.compile(r"/members/(\d+)")

# ------------------------------------------------------------------ roles
#
# A role is a short label - "Environment admin", "Operator", "Agent builder". Row
# and page text also carries names, emails, auth methods, statuses, avatar
# initials and activity lines ("Recipe started 2 days ago"), none of which is a
# role and all of which used to end up in the Role column. These filters exist so
# the column holds a role or nothing: blank is honest, raw row text is not.
ROLE_WORDS = ("admin", "developer", "operator", "analyst", "viewer", "member", "owner",
              "custom", "manager", "agent", "builder", "editor", "contributor",
              "designer", "architect", "workspace")
NOISE_RE = re.compile(
    r"^(sso|saml|oauth|password|active|inactive|invited|pending|disabled|deactivated|"
    r"logged\s*in.*|last\s*(active|login).*|never\s*logged.*|\d{1,2}/\d{1,2}/\d{2,4}.*|"
    r"[A-Z]{1,2}|yes|no|-|\u2014)$", re.I)
ENV_NAMES = ("development", "dev", "test", "staging", "stage", "uat", "production", "prod")
ACTIVITY_RE = re.compile(r"^(recent activity|activity|last (active|login|seen)|logged in|"
                         r"created|updated|invited|joined)\b", re.I)
# Lines that describe an event, not a role ("Custom oauth key deleted").
EVENT_RE = re.compile(r"\b(deleted|created|updated|started|stopped|added|removed|changed|"
                      r"edited|deployed|logged|signed|invited|key|token|oauth|connection|"
                      r"recipe|ago|\d{4})\b", re.I)


def _role_parts(cells, email, name=""):
    """Short text parts of a row that could be a role."""
    user = email.split("@")[0].lower()
    parts = []
    for c in cells:
        for part in re.split(r"\n|,\s*|\s{2,}|\s\|\s", c):
            part = part.strip(" |")
            if not part or "@" in part or len(part) > 60:
                continue
            low = part.lower()
            if low == user or (name and low == name.lower()) or NOISE_RE.match(part):
                continue
            parts.append(part)
    return parts


def guess_role(cells, email, name=""):
    """Role names only, from parts that contain a known role word. Never raw text."""
    known = [p for p in _role_parts(cells, email, name)
             if any(w in p.lower() for w in ROLE_WORDS) and not EVENT_RE.search(p)]
    return " | ".join(dict.fromkeys(known))


def role_for_env(cells, headers, email, name, env):
    """List page: when a column header names the environment, the role is that cell."""
    if not env or not headers:
        return None
    for i, h in enumerate(headers):
        if env.lower() in h.lower() and i < len(cells):
            parts = _role_parts([cells[i]], email, name)
            return " | ".join(dict.fromkeys(parts)) if parts else None
    return None


def _is_env_heading(s):
    s = s.rstrip(":").strip().lower()
    return s in ENV_NAMES or any(s in (f"{e} environment", f"{e} role", f"{e} roles")
                                 for e in ENV_NAMES)


def _role_like(s):
    return (any(w in s.lower() for w in ROLE_WORDS) and len(s) < 40 and "@" not in s
            and not NOISE_RE.match(s) and not EVENT_RE.search(s) and not _is_env_heading(s))


def role_from_detail_text(text, env):
    """Detail page: the role listed for `env`.

    Workato lays this out three ways depending on plan and page version - an
    environment heading with the role beneath it, "Development: Operator" on one
    line, or the role with the environment name beside it - so all three are
    tried. Without an environment match, the first role-like line before any
    activity section, preferring something more specific than plain "Member".
    Returns "" when nothing role-like is there.
    """
    lines = [l.strip() for l in text.splitlines() if l.strip()]
    low = [l.lower() for l in lines]
    env_l = (env or "").lower()
    if env_l:
        for i, l in enumerate(low):
            if _is_env_heading(l) and env_l in l:
                for j in range(i + 1, min(i + 6, len(lines))):
                    if _is_env_heading(low[j]) or ACTIVITY_RE.match(lines[j]):
                        break
                    if _role_like(lines[j]):
                        return lines[j]
        for l in lines:
            m = re.match(rf"^{re.escape(env_l)}\b\s*(environment)?\s*(role)?\s*[:\-\u2013]\s*(.+)$",
                         l, re.I)
            if m and _role_like(m.group(3).strip()):
                return m.group(3).strip()
        for i, l in enumerate(lines):
            if _role_like(l):
                window = low[max(0, i - 2):i] + low[i + 1:i + 3]
                if any(env_l == w.rstrip(":") or w.startswith(env_l) for w in window):
                    return l
    picked = []
    for l in lines:
        if ACTIVITY_RE.match(l):
            break
        if _role_like(l):
            picked.append(l)
    specific = [p for p in picked if p.lower() != "member"]
    return (specific or picked or [""])[0]


# Finds a usable "next page" control and marks it for the click. Covers aria
# labels, Next / > / >> text, rel=next, and numbered pagination (current + 1).
NEXT_PAGE_JS = r"""
() => {
  const isOff = el => el.disabled || el.getAttribute('aria-disabled') === 'true' ||
                      /disabled/i.test(el.className || '') ||
                      el.closest('[aria-disabled="true"], .disabled, [disabled]');
  const cands = Array.from(document.querySelectorAll('button, a, [role=button]'));
  const txt = el => ((el.innerText || el.textContent || '') + ' ' +
                     (el.getAttribute('aria-label') || '') + ' ' +
                     (el.getAttribute('title') || '')).trim();
  document.querySelectorAll('[data-uar-next]').forEach(e => e.removeAttribute('data-uar-next'));
  let next = cands.find(el => /^(next( page)?|›|»|>|→)$/i.test(txt(el)) ||
                              /next/i.test(el.getAttribute('aria-label') || '') ||
                              el.getAttribute('rel') === 'next');
  if (!next) {
    const cur = cands.find(el => el.getAttribute('aria-current') === 'page' ||
                                 /(active|current|selected)/i.test(el.className || ''));
    const n = cur ? parseInt(txt(cur), 10) : NaN;
    if (!isNaN(n)) next = cands.find(el => txt(el) === String(n + 1));
  }
  if (!next) return {found: false};
  if (isOff(next)) return {found: true, enabled: false};
  next.setAttribute('data-uar-next', '1');
  return {found: true, enabled: true};
}
"""


# ------------------------------------------------------------ terminations

def load_terminations(path):
    """Read the Workday "SOX Audit - Terminations" export, .csv or .xlsx.

    Returns (rows including the header, {email: termination date}). The header
    is the first row with an Email column, so a title row above it is fine.
    """
    p = Path(path).expanduser()
    rows = []
    if p.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        ws = load_workbook(p, read_only=True, data_only=True).worksheets[0]
        for r in ws.iter_rows(values_only=True):
            rows.append(["" if v is None else
                         (v.strftime("%m/%d/%y") if hasattr(v, "strftime") else str(v))
                         for v in r])
    else:
        with open(p, newline="", encoding="utf-8-sig") as f:
            rows = [list(r) for r in csv.reader(f)]
    hi = next((i for i, r in enumerate(rows)
               if any("email" in (c or "").lower() for c in r)), None)
    if hi is None:
        raise ValueError(f"{p.name}: no header row with an Email column")
    header = rows[hi]
    data = [r for r in rows[hi + 1:] if any((c or "").strip() for c in r)]
    ei = next(i for i, c in enumerate(header) if "email" in (c or "").lower())
    ti = next((i for i, c in enumerate(header)
               if "termination date" in (c or "").lower()), None)
    term = {}
    for r in data:
        if ei < len(r) and r[ei].strip():
            date = r[ti].strip() if ti is not None and ti < len(r) else ""
            term[r[ei].strip().lower()] = date or "Terminated"
    return [header] + data, term


def parse_pages(value):
    """Further access pages: "Label|URL" items, separated by ";" or newlines, or
    a YAML list of {label, url} / "Label|URL". Items without an http(s) URL are
    dropped - a page nobody can open is not a page to evidence."""
    items = value if isinstance(value, list) else re.split(r"[;\n]", str(value or ""))
    out = []
    for it in items:
        if isinstance(it, dict):
            label, url = str(it.get("label", "")).strip(), str(it.get("url", "")).strip()
        else:
            label, _, url = str(it).partition("|")
            label, url = label.strip(), url.strip()
            if not url and label.startswith("http"):
                label, url = "", label
        if re.match(r"^https?://", url):
            out.append({"label": label or url, "url": url})
    return out


class UarCapture:
    def __init__(self, page, out_dir, settle, workspace, display, max_parts,
                 max_pages=20):
        self.page = page
        self.out_dir = out_dir
        self.settle = settle
        self.workspace = workspace
        self.display = display
        self.max_parts = max_parts
        self.max_pages = max_pages
        self.seq = 0
        self.manifest = []
        self.warnings = []
        self.rejected = []
        self.no_access_control = False

    def warn(self, msg):
        log(f"  !! {msg}")
        self.warnings.append(msg)

    def shoot(self, kind, desc, url, parts=1, tab="before"):
        """Whole-screen capture, scrolling for extra parts on long listings."""
        self.seq += 1
        self.page.bring_to_front()
        activate_app(APP_NAME)
        time.sleep(1.0)
        ts = now_stamp()
        fname = (f"{self.seq:02d}_uar_{kind}_{ts.strftime('%Y%m%d-%H%M%S')}"
                 f"{'' if parts == 1 else f'_p{parts}'}.png")
        path = self.out_dir / fname
        grab_screen(path, self.display)
        log(f"  captured -> {fname}")
        self.manifest.append({
            "seq": self.seq, "part": parts, "kind": kind, "description": desc,
            "url": url, "captured": ts.isoformat(), "file": str(path), "tab": tab,
        })

    @staticmethod
    def roster_verdict(data, requested, landed):
        """Decide whether a page is really a collaborator roster.

        Returns (is_roster, reason). Containing an email address is not enough -
        Workato's admin error-alerts screen contains one, in a notification-recipient
        field, and accepting it produced a "user access review" listing one person
        who was not a collaborator at all. A roster has to show its structure.
        """
        # A redirect off the requested path means Workato sent us to its default
        # admin sub-page rather than the page asked for.
        req_path = requested.split("?")[0].rstrip("/")
        landed_path = landed.split("?")[0].split("#")[0].rstrip("/")
        redirected = not landed_path.startswith(req_path)

        rows = data.get("rows", [])
        if not rows:
            return False, "no rows with an email address"

        body = data.get("bodyText", "")
        words = [w for w in ROSTER_WORDS if w in body]
        headers = [h.lower() for h in data.get("headers", [])]
        header_hits = [h for h in headers if any(k in h for k in ROSTER_HEADERS)]

        if redirected and not header_hits:
            return False, (f"redirected to {landed_path} and it has no roster columns "
                           f"- this is not the page that was asked for")

        # A real table with roster-looking columns is conclusive, even for a
        # one-person workspace.
        if data.get("fromTable", 0) >= 1 and header_hits:
            return True, f"table with columns {header_hits}"

        # No table: several distinct people plus roster wording on the page.
        if len(rows) >= 3 and words:
            return True, f"{len(rows)} people listed, page mentions {words[:3]}"

        return False, (f"only {len(rows)} email-bearing row(s), no roster columns"
                       f"{' and no roster wording' if not words else ''}"
                       f" - looks like a settings page, not a roster")

    def find_members_page(self):
        """Open the first candidate URL that actually renders a collaborator roster."""
        rejected = []
        for url in MEMBER_URLS:
            log(f"  trying {url}")
            try:
                self.page.goto(url, wait_until="domcontentloaded")
            except Exception as exc:
                log(f"    could not open ({type(exc).__name__})")
                rejected.append(f"{url}: could not open")
                continue
            try:
                self.page.wait_for_load_state("networkidle", timeout=20000)
            except Exception:
                pass
            time.sleep(self.settle)

            landed = self.page.url
            if "/users/sign_in" in landed:
                self.warn("the session dropped while looking for the collaborator page")
                return None

            if "/members/access_control" in landed and self._missing_page():
                # The collaborator page, and Workato saying this account may not
                # see it. Recorded so the failure can name the real cause.
                self.no_access_control = True
                log("    rejected: Workato says this page does not exist for this account")
                rejected.append(f"{url} -> {landed}: page does not exist for this account")
                continue

            data = self.page.evaluate(EXTRACT_JS)
            ok, why = self.roster_verdict(data, url, landed)
            if ok:
                log(f"    accepted: {why}")
                if landed.rstrip("/") != url.rstrip("/"):
                    log(f"    (landed on {landed})")
                return landed
            log(f"    rejected: {why}")
            rejected.append(f"{url} -> {landed}: {why}")

        self.rejected = rejected
        return None

    def _missing_page(self):
        try:
            title = (self.page.title() or "").lower().replace("\u2019", "'")
        except Exception:
            return False
        return "page doesn't exist" in title or "page does not exist" in title

    def not_found_reason(self):
        """Why no roster was found, as something a person can act on."""
        if not getattr(self, "no_access_control", False):
            return ("Open the browser session, navigate to the collaborators list by hand, "
                    "and add that URL to MEMBER_URLS.")
        where = ("" if (self.workspace or "").lower() in ("development", "dev") else
                 f" Collaborators are managed only in the Development environment, and this "
                 f"run was for {self.workspace!r}; the roster is shared, so prepare the "
                 f"review with Development selected.")
        return ("Workato answered \"That page doesn't exist\" for Workspace admin > Access "
                "control > Collaborators. The account the runner signs in with is not "
                "allowed to manage collaborators in this workspace - its Workspace admin "
                "shows Settings only." + where + " Ask a Workato admin to give that account "
                "a Development role that includes collaborator management (or view access to "
                "it), or sign in to the browser session as someone who has it, then prepare "
                "the review again.")

    def goto(self, url):
        self.page.goto(url, wait_until="domcontentloaded")
        try:
            self.page.wait_for_load_state("networkidle", timeout=20000)
        except Exception:
            pass
        time.sleep(self.settle)

    def collect(self, url, label="Workato collaborators"):
        """Read every page of the listing.

        Each page is scrolled for rows rendered only on demand, then "next" is
        followed until it is missing or disabled, a page adds nobody new, or
        max_pages is reached - the last raised as a warning, because stopping
        early silently would under-count the population.
        """
        by_email, headers, parts = {}, [], 0
        for page_no in range(1, self.max_pages + 1):
            before = len(by_email)
            tag = label if page_no == 1 else f"{label} - page {page_no}"
            got, hdrs, n = self._collect_screen(url, tag, page_no)
            headers = headers or hdrs
            for row in got:
                by_email.setdefault(row["email"].lower(), row)
            parts += n
            nxt = self.page.evaluate(NEXT_PAGE_JS) or {}
            if not nxt.get("found") or not nxt.get("enabled"):
                break
            if page_no > 1 and len(by_email) == before:
                log("  next page added nobody new - stopping")
                break
            if page_no == self.max_pages:
                self.warn(f"'{label}' has more than {self.max_pages} pages; the rest were "
                          f"not read - the user list is incomplete")
                break
            try:
                self.page.click("[data-uar-next]", timeout=10000)
            except Exception as exc:
                self.warn(f"could not open page {page_no + 1} of '{label}' "
                          f"({str(exc).splitlines()[0]}) - the user list may be incomplete")
                break
            try:
                self.page.wait_for_load_state("networkidle", timeout=15000)
            except Exception:
                pass
            time.sleep(max(2.0, self.settle))
            self.page.evaluate("() => { const el = document.querySelector('[data-sox-scroll]');"
                               " if (el) el.scrollTop = 0; window.scrollTo(0, 0); }")
            log(f"  page {page_no + 1} of '{label}'")
        return list(by_email.values()), headers, parts

    def _collect_screen(self, url, label, page_no):
        """One page of a listing, scrolled through, one screenshot per screen."""
        by_email = {}
        headers = []
        parts = 0

        info = self.page.evaluate(SCROLL_JS)
        planned = 1
        if info:
            remaining = info["total"] - info["height"] - info["top"]
            if remaining > info["height"] * 0.25:
                planned = min(self.max_parts,
                              1 + math.ceil(remaining / max(1, info["height"] * 0.9)))

        for part in range(1, planned + 1):
            if part > 1:
                before = self.page.evaluate(
                    "() => (document.querySelector('[data-sox-scroll]') "
                    "|| document.scrollingElement).scrollTop")
                after = self.page.evaluate(SCROLL_BY_JS, int(info["height"] * 0.9))
                if after - before < 100:
                    break
                time.sleep(1.0)

            data = self.page.evaluate(EXTRACT_JS)
            if data["headers"] and not headers:
                headers = data["headers"]
            for row in data["rows"]:
                by_email.setdefault(row["email"].lower(), row)
            parts += 1
            self.shoot("collaborators", f"{label} - {url}", url,
                       parts=part if page_no == 1 else f"{page_no}.{part}")

        return list(by_email.values()), headers, parts

    def to_users(self, rows, headers, source="Workato collaborators"):
        users = []
        for row in rows:
            cells = row["cells"]
            email = row["email"]
            name = name_from_row(cells, headers, email)
            role = role_for_env(cells, headers, email, name, self.workspace)
            if role is None:
                role = (field_from_row(cells, headers, ("role", "permission", "access"))
                        or guess_role(cells, email, name))
            link = MEMBER_LINK_RE.search(row.get("href") or "")
            raw = status_from_row(cells, headers)
            active, bucket = classify(raw)
            if bucket == "unknown":
                self.warn(f"{email}: status {raw!r} not recognised - classify this by hand "
                          f"before signing off (counted as neither active nor inactive)")
            users.append({
                "name": name,
                "email": email,
                "role": role,
                "status_raw": raw,
                "status": bucket,
                "active": active,
                "last_activity": field_from_row(cells, headers,
                                                ("last", "activity", "seen", "login")),
                "member_id": link.group(1) if link else "",
                "source": source,
                "terminated": "",
                "source_row": row["text"][:500],
            })
        users.sort(key=lambda u: (u["status"] != "active", u["email"].lower()))
        return users

    def details(self, users, max_detail):
        """Open each collaborator's own page for the role in this environment.

        The roster shows one role; the detail page shows one per environment, and
        this review is of one environment. The detail page wins when it names a
        role. When it does not, the roster's role stays and a warning says so -
        the person may hold no role here, which a reviewer must confirm.
        """
        linked = [u for u in users if u.get("member_id")]
        if not linked:
            log("  rows do not link to collaborator pages - roles are as the list shows them")
            return
        if len(linked) > max_detail:
            self.warn(f"{len(linked)} collaborators, detail pages read for the first "
                      f"{max_detail} only (--max-detail); the rest show the list's role")
        for u in linked[:max_detail]:
            url = f"{BASE}/members/{u['member_id']}"
            log(f"  collaborator {u['email']}")
            try:
                self.goto(url)
            except Exception as exc:
                self.warn(f"{u['email']}: detail page did not open ({type(exc).__name__})")
                continue
            if "/users/sign_in" in self.page.url:
                self.warn("the session dropped while reading collaborator pages")
                return
            role = role_from_detail_text(page_text(self.page), self.workspace)
            if role:
                u["role"] = role
            else:
                self.warn(f"{u['email']}: no role for '{self.workspace}' on their page - "
                          f"they may hold none in this environment; check the screenshot")
            self.shoot("collaborator", f"Collaborator - {u['name']} ({u['email']})", url)

    def extra_page(self, pg):
        """A further access page (Workato Agentic users, say), listed as-is."""
        log(f"Page: {pg['label']} - {pg['url']}")
        try:
            self.goto(pg["url"])
        except Exception as exc:
            self.warn(f"'{pg['label']}' did not open ({type(exc).__name__})")
            return []
        if "/users/sign_in" in self.page.url or "login" in self.page.url.lower():
            self.warn(f"'{pg['label']}' sent the browser to a sign-in page - check the URL "
                      f"and this account's access to it")
            return []
        rows, headers, _ = self.collect(pg["url"], pg["label"])
        if not rows:
            self.warn(f"no users read from '{pg['label']}' - fill its rows in from the "
                      f"screenshots; this is not evidence it has no users")
        return self.to_users(rows, headers, source=pg["label"])


def apply_terminations(users, term, warn):
    """Mark each user who appears on the terminations list. A terminated person
    who still holds access is the finding a UAR exists to surface, so each one is
    a warning as well as a column."""
    for u in users:
        date = term.get(u["email"].lower(), "")
        u["terminated"] = date
        if date and u["active"] is not False:
            warn(f"{u['email']} was terminated {date} but is still listed in "
                 f"{u['source']} ({u['status_raw'] or u['status']}) - access needs review")


def write_outputs(out_dir, users, meta):
    (out_dir / "users.json").write_text(json.dumps(
        {**meta, "users": users}, indent=2))

    cols = ["name", "email", "role", "status", "status_raw", "active", "terminated",
            "last_activity", "source", "member_id"]
    with (out_dir / "users.csv").open("w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=cols, extrasaction="ignore")
        w.writeheader()
        for u in users:
            w.writerow(u)
    log(f"  wrote users.csv and users.json ({len(users)} users)")


BOLD = Font(bold=True)
HDR_FILL = PatternFill("solid", fgColor="D9E1F2")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
TERM_FILL = PatternFill("solid", fgColor="F8CBAD")
UNKNOWN_FILL = PatternFill("solid", fgColor="FCE4D6")
ROLES_DOC = ("https://docs.workato.com/en/user-accounts-and-teams/role-based-access/"
             "new-model/system-environment-roles")

CHECKLIST = [
    '1. In tab "1. Screenshot and Parameters" - Was a screenshot taken of the parameters '
    'used to generate the listing which includes the date/time the listing was generated?',
    '2. In tab "2. User Listing" - Was the listing reviewed and tick-marked for all users '
    'regardless of whether they are appropriate or not?',
    '3. If inappropriate access was identified in step 2; Does tab "3. Screenshots and '
    'Params Re-Run" contain a screenshot of the parameters used to generate the validation '
    'listing which includes the date/time the listing was generated?',
    '4. If inappropriate access was identified in step 2; Was a look back performed and '
    'documented in "Look Back Evidence Tab" to confirm each high risk user deemed '
    'inappropriate during the QAR did not perform any activities which could be deemed '
    'inappropriate?',
    '5. If inappropriate access was identified in step 3; Does tab "User Listing Re-Run" '
    'contain the full re-ran listing and positive confirmation of all users appropriateness?',
    '6. Is there sign-off from both the primary reviewer and the secondary reviewer?',
]


def _place_images(ws, shots, row, img_dir, max_width):
    for c in shots:
        ws.cell(row=row, column=1, value=c["description"]).font = BOLD
        ts = datetime.fromisoformat(c["captured"]).strftime("%Y-%m-%d %H:%M:%S %Z")
        ws.cell(row=row + 1, column=1, value=f"{c['url']}   ·   captured {ts}")
        img = XLImage(workbook_copy(c["file"], img_dir))
        if img.width > max_width:
            k = max_width / img.width
            img.width, img.height = int(img.width * k), int(img.height * k)
        ws.add_image(img, f"A{row + 2}")
        row += 2 + math.ceil(img.height / PX_PER_ROW) + 3
    return row


def build_workbook(manifest_path, out_path, max_width=1400):
    """The UAR template: checklist, screenshots before, user list, Workday
    terminations, screenshots after."""
    data = json.loads(Path(manifest_path).read_text())
    users = data["users"]
    period, ws_name = data.get("period", ""), data.get("workspace", "")
    checked = bool(data.get("terminations_file"))
    img_dir = Path(manifest_path).parent / "_workbook_images"
    listed = (data.get("captured") or "").replace("T", " ")[:19]

    wb = Workbook()
    wb.remove(wb.active)

    # 1. Review Checklist ------------------------------------------------------
    ws = wb.create_sheet("Review Checklist")
    ws.column_dimensions["A"].width = 110
    ws.column_dimensions["B"].width = 45
    head = [("SYSTEM", f"Workato ({ws_name} environment)" if ws_name else "Workato"),
            ("Period", period), ("Listing taken", listed),
            ("Source page", data.get("source_url", "")),
            ("Users listed", len(users)),
            ("Active", sum(1 for u in users if u["active"] is True)),
            ("Inactive / suspended", sum(1 for u in users if u["status"] == "inactive")),
            ("Pending invitation", sum(1 for u in users if u["status"] == "pending")),
            ("Status not recognised", sum(1 for u in users if u["active"] is None)),
            ("Terminated, still listed",
             sum(1 for u in users if u.get("terminated")) if checked else "not checked")]
    for i, (k, v) in enumerate(head, start=1):
        ws.cell(row=i, column=1, value=k).font = BOLD
        ws.cell(row=i, column=2, value=v)
    r = len(head) + 2
    ws.cell(row=r, column=1, value="Checklist for QAR Listings").font = BOLD
    r += 1
    for q in CHECKLIST:
        ws.cell(row=r, column=1, value=q).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=r, column=2).fill = INPUT_FILL
        r += 1
    ws.cell(row=r, column=1, value="Is my QAR Ready?").font = BOLD
    ws.cell(row=r, column=2).fill = INPUT_FILL
    ws.cell(row=r + 2, column=1, value="Primary Reviewer Sign Off, Date, and Title")
    ws.cell(row=r + 2, column=2).fill = INPUT_FILL
    ws.cell(row=r + 3, column=1, value="Secondary Reviewer Sign Off, Date, and Title")
    ws.cell(row=r + 3, column=2).fill = INPUT_FILL
    note = ws.cell(row=r + 5, column=1, value=(
        f"Point-in-time listing taken {listed} and labelled {period}. Workato publishes no "
        "historical collaborator list, so this evidences access as it stood when the "
        "capture ran - not on the last day of the period."))
    note.alignment = Alignment(wrap_text=True)
    ws.row_dimensions[r + 5].height = 44
    if data.get("warnings"):
        ws.cell(row=r + 7, column=1, value="Capture warnings - resolve before signing off").font = BOLD
        for i, w in enumerate(data["warnings"], start=r + 8):
            ws.cell(row=i, column=1, value=w).alignment = Alignment(wrap_text=True)

    # 2. Parameter Screenshot - Before ---------------------------------------
    ws = wb.create_sheet("Parameter Screenshot - Before")
    ws.column_dimensions["A"].width = 200
    ws["A1"] = (f"Workato collaborators and roles - {ws_name} - captured {listed} "
                "(the on-screen clock in each image is the date/time evidence)")
    ws["A1"].font = BOLD
    before = [c for c in data.get("captures", []) if c.get("tab", "before") == "before"]
    if before:
        _place_images(ws, before, 3, img_dir, max_width)
    else:
        ws["A3"] = "No screenshots were captured."

    # 3. User List -------------------------------------------------------------
    ws = wb.create_sheet("User List")
    cols = [("Name", 28), ("Email", 34), (f"Access Type ({ws_name})" if ws_name else "Access Type", 30),
            ("Status (as shown)", 18), ("Verdict", 14), ("Is the user terminated?", 24),
            ("Access Appropriate?", 22), ("Access Reviewed By", 24), ("Reviewed On", 16),
            ("Source (platform / page)", 30)]
    for i, (h, w) in enumerate(cols, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font, c.fill = BOLD, HDR_FILL
        ws.column_dimensions[c.column_letter].width = w
    ws.freeze_panes = "A2"
    for r, u in enumerate(users, start=2):
        verdict = {"active": "Active", "inactive": "Inactive", "pending": "Pending"}.get(
            u["status"], "Unrecognised")
        term = (u.get("terminated") or "Active User") if checked else "not checked"
        vals = [u["name"], u["email"], u.get("role") or "", u.get("status_raw") or "",
                verdict, term, None, None, None, u.get("source", "")]
        for i, v in enumerate(vals, start=1):
            ws.cell(row=r, column=i, value=v)
        for i in (7, 8, 9):
            ws.cell(row=r, column=i).fill = INPUT_FILL
        if u["active"] is None:
            ws.cell(row=r, column=5).fill = UNKNOWN_FILL
        if u.get("terminated"):
            ws.cell(row=r, column=6).fill = TERM_FILL
    end = len(users) + 3
    if not users:
        ws.cell(row=2, column=1, value=("No users could be read - fill this in from the "
                                        "Before screenshots. This is not evidence of no users."))
    ws.cell(row=end, column=1, value=(
        "Access Type = the role in this environment, from each collaborator's own page where "
        "it could be read, else as the list shows it. Columns G-I are for the reviewer. A "
        "user listed on more than one page (Workato and Workato Agentic) appears once per "
        "page.")).font = Font(italic=True)

    # 4. Workday termination List --------------------------------------------
    ws = wb.create_sheet("Workday termination List")
    ws["A1"] = "SOX Audit - Terminations"
    ws["A1"].font = BOLD
    rows = data.get("termination_rows") or []
    if rows:
        ws["A2"] = f"Source: {data.get('terminations_file', '')}"
        for ri, row in enumerate(rows, start=4):
            for ci, v in enumerate(row, start=1):
                c = ws.cell(row=ri, column=ci, value=v)
                if ri == 4:
                    c.font, c.fill = BOLD, HDR_FILL
    else:
        ws["A2"] = ("Not checked: no terminations list was supplied. Paste the Workday "
                    "'SOX Audit - Terminations' export here and fill column F of the User "
                    "List, or set uar_terminations for Workato and prepare again to have it "
                    "cross-checked automatically.")

    # 5. Parameter Screenshot - After ----------------------------------------
    ws = wb.create_sheet("Parameter Screenshot - After")
    ws.column_dimensions["A"].width = 200
    after = [c for c in data.get("captures", []) if c.get("tab") == "after"]
    if after:
        r = _place_images(ws, after, 1, img_dir, max_width)
    else:
        ws["A1"] = ("NA - no inappropriate access identified; no re-run required. "
                    "(Replace with re-run screenshots if applicable.)")
        r = 3
    ws.cell(row=r + 1, column=1, value=f"Additional Details: {ROLES_DOC}")

    wb.save(out_path)
    log(f"Workbook written: {out_path}")


def run(args, out_dir):
    with sync_playwright() as p:
        ctx = launch_browser(p, args.profile, args.browser)
        page = ctx.pages[0] if ctx.pages else ctx.new_page()
        page.set_default_timeout(60000)

        # Same sign-in path as the change-management capture, including the
        # configured credentials and the fallback to a person for SSO/MFA.
        from workato.sox_capture import wait_for_login

        page = wait_for_login(ctx, page, username=args.username,
                              password=args.password, mode=args.login)
        if page is None:
            log("Login did not complete. Exiting.")
            try:
                ctx.close()
            except Exception:
                pass
            sys.exit(1)
        page.set_default_timeout(60000)
        page.bring_to_front()
        log(f"Logged in. Current URL: {page.url}")

        if args.workspace:
            phase("checking_workspace", args.workspace)
            ensure_workspace(page, args.workspace)

        phase("capturing", "collaborators")
        cap = UarCapture(page, out_dir, args.settle, args.workspace, args.display,
                         args.max_parts, args.max_pages)

        url = cap.find_members_page()
        if not url:
            reason = cap.not_found_reason()
            cap.warn(reason)
            cap.shoot("not-found", "No collaborator page could be opened", page.url)
            manifest = out_dir / "manifest.json"
            manifest.write_text(json.dumps({
                "period": args.period, "workspace": args.workspace,
                "captured": now_stamp().isoformat(), "source_url": "",
                "users": [], "captures": cap.manifest, "pages": [],
                "rejected": cap.rejected,
                "warnings": cap.warnings + [
                    "No collaborator roster could be found at any known URL. No user "
                    "list was produced - this is NOT evidence that there are no users."],
            }, indent=2))
            ctx.close()
            raise SystemExit(
                "No collaborator roster could be found. Each candidate was opened and "
                "rejected:\n  " + "\n  ".join(cap.rejected or list(MEMBER_URLS))
                + "\n\nNo user list was produced. This is NOT evidence that the workspace "
                  "has no users.\n" + reason)

        log(f"Collaborator page: {url}")
        rows, headers, parts = cap.collect(url)
        log(f"  {len(rows)} collaborator row(s) across {parts} screen(s)")
        users = cap.to_users(rows, headers)

        if not users:
            cap.warn("the page opened but no collaborator rows could be read - the listing "
                     "may be rendered in a way this tool does not understand. Check the "
                     "screenshots; do not read this as an empty user list.")
        elif not args.no_detail:
            phase("capturing", "collaborator roles")
            cap.details(users, args.max_detail)

        pages = [{"label": "Workato collaborators", "url": url}]
        for pg in args.pages:
            phase("capturing", pg["label"])
            users += cap.extra_page(pg)
            pages.append(pg)

        if args.term_map is not None:
            apply_terminations(users, args.term_map, cap.warn)

        meta = {
            "period": args.period,
            "workspace": args.workspace,
            "captured": now_stamp().isoformat(),
            "source_url": url,
            "pages": pages,
            "terminations_file": Path(args.terminations).name if args.terminations else "",
        }
        write_outputs(out_dir, users, meta)

        manifest = out_dir / "manifest.json"
        manifest.write_text(json.dumps({
            **meta, "users": users, "captures": cap.manifest, "warnings": cap.warnings,
            # The Workday rows go to the workbook only. users.json is the listing
            # of who has access; it has no reason to carry who left the company.
            "termination_rows": args.term_rows,
        }, indent=2))
        log(f"Manifest written: {manifest} ({len(cap.manifest)} screenshots)")

        active = sum(1 for u in users if u["active"] is True)
        inactive = sum(1 for u in users if u["active"] is False)
        unknown = sum(1 for u in users if u["active"] is None)
        log(f"Users: {len(users)} total · {active} active · {inactive} inactive/pending "
            f"· {unknown} unrecognised")
        if args.term_map is not None:
            gone = sum(1 for u in users if u["terminated"])
            log(f"Terminations: {len(args.term_map)} on the Workday list, {gone} still listed here")

        if cap.warnings:
            print("\n==== WARNINGS - resolve these before signing off ====")
            for w in cap.warnings:
                print("  " + w)
            print()
        ctx.close()
        return manifest


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--period", default="", help='review period label, e.g. "Q3 FY26"')
    ap.add_argument("--out", default=None, help="output folder")
    ap.add_argument("--profile", default="~/.workato_sox_browser_profile")
    ap.add_argument("--browser", choices=["chromium", "chrome"], default="chromium")
    ap.add_argument("--settle", type=float, default=4.0)
    ap.add_argument("--workspace", default="Production")
    ap.add_argument("--display", type=int, default=1)
    ap.add_argument("--max-parts", type=int, default=12,
                    help="max scrolled screens per page of the user list")
    ap.add_argument("--max-pages", type=int, default=20,
                    help="max pages followed on a paginated user list")
    ap.add_argument("--no-detail", action="store_true",
                    help="skip each collaborator's own page (roles as the list shows them)")
    ap.add_argument("--max-detail", type=int, default=60,
                    help="max collaborator pages read for their environment role")
    ap.add_argument("--terminations", default="",
                    help="Workday 'SOX Audit - Terminations' export (.csv or .xlsx); "
                         "default: the environment's uar_terminations setting")
    ap.add_argument("--pages", default="",
                    help='further access pages, "Label|URL; Label|URL"; '
                         "default: the environment's uar_pages setting")
    ap.add_argument("--capture-only", action="store_true")
    ap.add_argument("--build-only", action="store_true")
    ap.add_argument("--no-pause", action="store_true")
    ap.add_argument("--env", default="workato")
    ap.add_argument("--username", default="")
    args = ap.parse_args()

    if not args.period:
        args.period = f"Q{(datetime.now().month - 1) // 3 + 1} {datetime.now().year}"

    args.password = ""
    args.login = "interactive"
    cfg = {}
    try:
        cfg = environments.find(args.env) or {}
        if cfg and cfg.get("enabled"):
            args.username = args.username or cfg.get("username", "")
            args.password = cfg.get("password", "")
            args.login = cfg.get("login", "interactive")
            if args.workspace == ap.get_default("workspace") and cfg.get("workspace"):
                args.workspace = cfg["workspace"]
    except Exception as exc:
        log(f"!! environments.yaml could not be read ({exc}); signing in by hand.")

    # A run from the page cannot pass these, so the environment's config (Doppler
    # or the YAML) supplies them; a flag on the command line wins.
    args.pages = parse_pages(args.pages or cfg.get("uar_pages", ""))
    args.terminations = str(args.terminations or cfg.get("uar_terminations", "") or "").strip()
    args.term_rows, args.term_map = [], None
    if args.terminations:
        try:
            args.term_rows, args.term_map = load_terminations(args.terminations)
            log(f"Terminations list: {Path(args.terminations).name}, "
                f"{len(args.term_map)} email(s)")
        except Exception as exc:
            # Not fatal: the roster is still worth capturing. But the check did not
            # happen, and the workbook must say "not checked", not "Active User".
            log(f"!! terminations list {args.terminations!r} could not be read ({exc}); "
                f"the review continues without the cross-check")
            args.terminations = ""

    out_dir = Path(args.out) if args.out else Path(
        f"uar_{slug(args.period)}_{datetime.now().strftime('%Y%m%d-%H%M%S')}")
    out_dir.mkdir(parents=True, exist_ok=True)
    xlsx = out_dir / f"Workato User Access Review - {args.period}.xlsx"

    if args.build_only:
        m = out_dir / "manifest.json"
        if not m.exists():
            sys.exit(f"No manifest at {m}")
        build_workbook(m, xlsx)
        return

    phase("starting", args.period)
    manifest = run(args, out_dir)
    if not args.capture_only:
        phase("building_workbook", str(xlsx))
        build_workbook(manifest, xlsx)
    phase("done", str(xlsx))
    log("Done.")


if __name__ == "__main__":
    main()
